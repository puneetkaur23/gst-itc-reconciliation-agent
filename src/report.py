import csv
import io
from collections import defaultdict
from datetime import datetime
from typing import Dict, List


class ReportGenerator:
    """Generate human-readable and machine-readable output files."""

    _EXCEPTION_CSV_COLUMNS = [
        "sr_no", "record_type", "invoice_number", "supplier_gstin", "supplier_name",
        "invoice_date", "taxable_value", "total_tax", "exception_code", "exception_reason",
        "action_required", "severity", "section_16_4_deadline", "days_to_deadline",
        "deadline_status", "rule_37a_triggered", "rule_37a_deadline",
    ]

    def generate_exception_csv(self, exceptions: List[Dict]) -> str:
        """Generate exception list as CSV string."""
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=self._EXCEPTION_CSV_COLUMNS, extrasaction="ignore")
        writer.writeheader()

        for i, exc in enumerate(exceptions, 1):
            rec = exc.get("record", {})
            rule37a = exc.get("rule_37a", {})
            row = {
                "sr_no": i,
                "record_type": exc.get("record_type", ""),
                "invoice_number": rec.get("invoice_number_raw", ""),
                "supplier_gstin": rec.get("supplier_gstin", ""),
                "supplier_name": rec.get("supplier_name", ""),
                "invoice_date": rec.get("invoice_date_raw", ""),
                "taxable_value": rec.get("taxable_value", ""),
                "total_tax": rec.get("total_tax", ""),
                "exception_code": exc.get("exception_code", ""),
                "exception_reason": exc.get("exception_reason", ""),
                "action_required": exc.get("action_required", ""),
                "severity": exc.get("severity", ""),
                "section_16_4_deadline": exc.get("section_16_4_deadline", ""),
                "days_to_deadline": exc.get("days_to_deadline", ""),
                "deadline_status": exc.get("deadline_status", ""),
                "rule_37a_triggered": rule37a.get("triggered", False),
                "rule_37a_deadline": rule37a.get("reversal_deadline", ""),
            }
            writer.writerow(row)

        return output.getvalue()

    def generate_summary_markdown(self, result: Dict) -> str:
        """Generate professional Markdown reconciliation report."""
        summary = result.get("summary", {})
        drc = result.get("drc01c_risk", {})
        exceptions = result.get("exceptions", [])
        ts = result.get("run_timestamp", "")

        lines = [
            "# GST ITC Reconciliation Report",
            "",
            f"**Generated:** {ts}",
            f"**Return Period:** {result.get('return_period', 'N/A')}",
            "",
            "---",
            "",
            "## Reconciliation Summary",
            "",
            "| Metric | Value |",
            "|--------|-------|",
            f"| Purchase Register Records | {summary.get('total_pr_records', 0)} |",
            f"| GSTR-2B Records | {summary.get('total_gstr2b_records', 0)} |",
            f"| AUTO MATCHED | {summary.get('auto_matched', 0)} ({summary.get('auto_match_pct', 0):.1f}%) |",
            f"| SUGGESTED MATCH (Review Required) | {summary.get('suggested_matched', 0)} ({summary.get('suggested_match_pct', 0):.1f}%) |",
            f"| UNMATCHED (Exceptions) | {summary.get('total_unmatched', 0)} ({summary.get('unmatched_pct', 0):.1f}%) |",
            f"| Overall Match Rate | **{summary.get('match_rate_pct', 0):.1f}%** |",
            f"| Eligible ITC (Matched) | Rs {summary.get('eligible_itc', 0):,.2f} |",
            f"| ITC at Risk (Unmatched) | Rs {summary.get('itc_at_risk', 0):,.2f} |",
            "",
            "---",
            "",
            "## DRC-01C Risk Assessment",
            "",
        ]

        if drc.get("risk"):
            lines += [
                "> **⚠️ HIGH RISK: DRC-01C Notice Likely**",
                "",
                f"- **Claimed ITC (GSTR-3B):** Rs {drc.get('claimed_itc', 0):,.2f}",
                f"- **Eligible ITC (from reconciliation):** Rs {drc.get('eligible_itc', 0):,.2f}",
                f"- **Excess Amount:** Rs {drc.get('excess_amount', 0):,.2f} ({drc.get('excess_percentage', 0):.2f}%)",
                f"- **Action Required:** {drc.get('action', '')}",
                "",
            ]
        else:
            lines += [
                "> **✅ No DRC-01C Risk Detected**",
                "",
                f"- Claimed ITC: Rs {drc.get('claimed_itc', 0):,.2f}",
                f"- Eligible ITC: Rs {drc.get('eligible_itc', 0):,.2f}",
                "",
            ]

        lines += [
            "---",
            "",
            "## Exception Breakdown",
            "",
        ]

        # Group exceptions by code
        by_code = defaultdict(list)
        for exc in exceptions:
            by_code[exc.get("exception_code", "UNKNOWN")].append(exc)

        severity_order = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3}
        for code in sorted(by_code.keys(), key=lambda c: severity_order.get(by_code[c][0].get("severity", "LOW"), 99)):
            group = by_code[code]
            sev = group[0].get("severity", "")
            sev_emoji = {"CRITICAL": "🔴", "HIGH": "🟠", "MEDIUM": "🟡", "LOW": "🟢"}.get(sev, "⚪")
            lines += [
                f"### {sev_emoji} {code} ({len(group)} records, Severity: {sev})",
                "",
                f"**Action:** {group[0].get('action_required', '')}",
                "",
                "| # | Invoice | Supplier GSTIN | Date | Taxable Value | Deadline Status |",
                "|---|---------|----------------|------|---------------|-----------------|",
            ]
            for j, exc in enumerate(group[:10], 1):
                rec = exc.get("record", {})
                lines.append(
                    f"| {j} | {rec.get('invoice_number_raw', '')} | "
                    f"{rec.get('supplier_gstin', '')} | "
                    f"{rec.get('invoice_date_raw', '')} | "
                    f"Rs {rec.get('taxable_value', 0):,.2f} | "
                    f"{exc.get('deadline_status', '')} |"
                )
            if len(group) > 10:
                lines.append(f"| ... | *{len(group) - 10} more records* | | | | |")
            lines.append("")

        lines += [
            "---",
            "",
            "## Audit Trail",
            "",
            "Full audit trail is available in udit_trail.json. Each matched pair includes:",
            "- Decision ID and timestamp",
            "- Raw and normalized invoice numbers",
            "- Dimension scores (number, date, taxable, tax) and weights",
            "- Normalization steps applied",
            "- Tolerance thresholds used",
            "",
        ]

        return "\n".join(lines)

    def generate_audit_trail(self, matched_pairs: List[Dict]) -> List[Dict]:
        """Generate a detailed audit entry for each matched pair."""
        trail = []
        for i, pair in enumerate(matched_pairs, 1):
            pr = pair["pr_record"]
            gstr = pair["gstr_record"]
            details = pair.get("match_details", {})

            # Detect what normalizations were applied to PR invoice number
            norm_applied = []
            raw = pr.get("invoice_number_raw", "")
            norm = pr.get("invoice_number_norm", "")
            if raw != norm:
                if raw.upper() != norm:
                    norm_applied.append("prefix_stripped")
                if any(c in raw for c in ["/", "-", "_"]):
                    norm_applied.append("separators_removed")
                if raw.lstrip("0") != raw:
                    norm_applied.append("leading_zeros_removed")
                if not norm_applied:
                    norm_applied.append("case_normalized")

            entry = {
                "decision_id": i,
                "timestamp": datetime.now().isoformat(),
                "match_class": pair.get("match_class", ""),
                "composite_score": round(pair.get("match_score", 0), 4),
                "pr_record": {
                    "invoice_number_raw": pr.get("invoice_number_raw"),
                    "invoice_number_norm": pr.get("invoice_number_norm"),
                    "invoice_date_raw": pr.get("invoice_date_raw"),
                    "invoice_date_norm": pr.get("invoice_date_norm"),
                    "supplier_gstin": pr.get("supplier_gstin"),
                    "supplier_name": pr.get("supplier_name"),
                    "taxable_value": pr.get("taxable_value"),
                    "total_tax": pr.get("total_tax"),
                    "document_type": pr.get("document_type"),
                },
                "gstr_record": {
                    "invoice_number_raw": gstr.get("invoice_number_raw"),
                    "invoice_number_norm": gstr.get("invoice_number_norm"),
                    "invoice_date_raw": gstr.get("invoice_date_raw"),
                    "invoice_date_norm": gstr.get("invoice_date_norm"),
                    "supplier_gstin": gstr.get("supplier_gstin"),
                    "supplier_name": gstr.get("supplier_name"),
                    "taxable_value": gstr.get("taxable_value"),
                    "total_tax": gstr.get("total_tax"),
                    "document_type": gstr.get("document_type"),
                },
                "dimension_scores": {
                    "number_score": round(details.get("number_score", 0), 4),
                    "date_score": round(details.get("date_score", 0), 4),
                    "taxable_score": round(details.get("taxable_score", 0), 4),
                    "tax_score": round(details.get("tax_score", 0), 4),
                },
                "weights_used": details.get("weights", {}),
                "normalization_applied": norm_applied if norm_applied else ["none"],
                "tolerance_used": {
                    "date_tolerance_days": 3,
                    "amount_tolerance_pct": "1%",
                    "amount_tolerance_abs": "Rs 5.00",
                },
            }
            trail.append(entry)

        return trail

    def generate_exception_xlsx(self, exceptions, filepath):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment
        from openpyxl.utils import get_column_letter
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Exceptions"
        
        headers = ['sr_no', 'record_type', 'invoice_number', 'supplier_gstin', 'supplier_name', 
                   'invoice_date', 'taxable_value', 'total_tax', 'exception_code', 
                   'exception_reason', 'action_required', 'severity', 'section_16_4_deadline',
                   'days_to_deadline', 'deadline_status', 'rule_37a_triggered', 'rule_37a_deadline']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        severity_fill = {
            'CRITICAL': PatternFill(start_color="C00000", end_color="C00000", fill_type="solid"),
            'HIGH': PatternFill(start_color="FF6600", end_color="FF6600", fill_type="solid"),
            'MEDIUM': PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid"),
            'LOW': PatternFill(start_color="00B050", end_color="00B050", fill_type="solid"),
        }
        
        for idx, exc in enumerate(exceptions, 1):
            row = idx + 1
            record = exc['record']
            data = {
                'sr_no': idx,
                'record_type': exc['record_type'],
                'invoice_number': record.get('invoice_number_raw', ''),
                'supplier_gstin': record.get('supplier_gstin', ''),
                'supplier_name': record.get('supplier_name', ''),
                'invoice_date': record.get('invoice_date_raw', ''),
                'taxable_value': record.get('taxable_value', 0),
                'total_tax': record.get('total_tax', 0),
                'exception_code': exc['exception_code'],
                'exception_reason': exc['exception_reason'],
                'action_required': exc['action_required'],
                'severity': exc['severity'],
                'section_16_4_deadline': str(exc.get('section_16_4_deadline', '')),
                'days_to_deadline': exc.get('days_to_deadline', ''),
                'deadline_status': exc.get('deadline_status', ''),
                'rule_37a_triggered': exc.get('rule_37a', {}).get('triggered', False),
                'rule_37a_deadline': str(exc.get('rule_37a', {}).get('reversal_deadline', '')),
            }
            
            for col, header in enumerate(headers, 1):
                value = data.get(header, '')
                cell = ws.cell(row=row, column=col, value=value)
                if header == 'severity' and value in severity_fill:
                    cell.fill = severity_fill[value]
                    cell.font = Font(bold=True, color="FFFFFF" if value in ['CRITICAL', 'HIGH'] else "000000")
                if header in ['taxable_value', 'total_tax']:
                    cell.number_format = '#,##0.00'
        
        for col in range(1, len(headers)+1):
            max_length = 0
            column = get_column_letter(col)
            for row in range(1, len(exceptions)+2):
                cell_value = ws.cell(row=row, column=col).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            ws.column_dimensions[column].width = min(max_length + 2, 50)
        
        wb.save(filepath)
        return filepath
